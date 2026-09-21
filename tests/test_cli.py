import json
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def run_cli(*arguments: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "office_design_builder.cli", *arguments],
        check=False,
        capture_output=True,
        text=True,
    )


def test_cli_help_lists_public_commands() -> None:
    result = run_cli("--help")

    assert result.returncode == 0
    assert "inspect" in result.stdout
    assert "validate" in result.stdout
    assert "build" in result.stdout
    assert "verify" in result.stdout


def test_validate_reports_stable_error_for_invalid_spec(tmp_path: Path) -> None:
    spec_path = tmp_path / "invalid.json"
    spec_path.write_text(json.dumps({"version": "1", "title": "Missing slides"}))

    result = run_cli("validate", str(spec_path))

    assert result.returncode != 0
    assert result.stderr.startswith("INVALID_SPEC:")
    assert str(spec_path) in result.stderr


def test_inspect_reports_stable_error_for_missing_input(tmp_path: Path) -> None:
    missing = tmp_path / "missing.png"
    output = tmp_path / "fingerprint.json"

    result = run_cli("inspect", str(missing), "--output", str(output))

    assert result.returncode != 0
    assert result.stderr.startswith("INPUT_NOT_FOUND:")
    assert str(missing) in result.stderr


def test_inspect_reports_stable_error_for_corrupt_png(tmp_path: Path) -> None:
    reference = tmp_path / "broken.png"
    output = tmp_path / "fingerprint.json"
    reference.write_bytes(b"not a png")

    result = run_cli("inspect", str(reference), "--output", str(output))

    assert result.returncode == 2
    assert result.stderr.startswith("INVALID_REFERENCE:")
    assert str(reference) in result.stderr
    assert "Traceback" not in result.stderr
    assert not output.exists()


def test_inspect_reports_stable_error_for_corrupt_pptx(tmp_path: Path) -> None:
    reference = tmp_path / "broken.pptx"
    output = tmp_path / "fingerprint.json"
    reference.write_bytes(b"not a pptx")

    result = run_cli("inspect", str(reference), "--output", str(output))

    assert result.returncode == 2
    assert result.stderr.startswith("INVALID_REFERENCE:")
    assert str(reference) in result.stderr
    assert "Traceback" not in result.stderr
    assert not output.exists()


def test_inspect_rejects_empty_pptx(tmp_path: Path) -> None:
    reference = tmp_path / "empty.pptx"
    output = tmp_path / "fingerprint.json"
    Presentation().save(str(reference))

    result = run_cli("inspect", str(reference), "--output", str(output))

    assert result.returncode == 2
    assert result.stderr.startswith("INVALID_REFERENCE:")
    assert "contains no slides" in result.stderr
    assert not output.exists()


def test_inspect_rejects_unsupported_extension(tmp_path: Path) -> None:
    reference = tmp_path / "reference.txt"
    output = tmp_path / "fingerprint.json"
    reference.write_text("not a supported reference")

    result = run_cli("inspect", str(reference), "--output", str(output))

    assert result.returncode == 2
    assert result.stderr.startswith("INVALID_REFERENCE:")
    assert "unsupported extension: .txt" in result.stderr
    assert "Traceback" not in result.stderr
    assert not output.exists()


def test_inspect_reports_stable_error_when_output_parent_is_missing(tmp_path: Path) -> None:
    reference = tmp_path / "reference.png"
    output = tmp_path / "missing" / "fingerprint.json"
    Image.new("RGB", (2, 2), color=(17, 34, 51)).save(reference)

    result = run_cli("inspect", str(reference), "--output", str(output))

    assert result.returncode == 2
    assert result.stderr.startswith("OUTPUT_WRITE_FAILED:")
    assert str(output) in result.stderr
    assert "Traceback" not in result.stderr
    assert not output.exists()


def test_inspect_writes_fingerprint_json(tmp_path: Path) -> None:
    reference_path = tmp_path / "reference.png"
    output_path = tmp_path / "fingerprint.json"
    Image.new("RGB", (4, 2), color=(17, 34, 51)).save(reference_path)

    result = run_cli("inspect", str(reference_path), "--output", str(output_path))

    assert result.returncode == 0, result.stderr
    payload = json.loads(output_path.read_text())
    assert payload["canvas"]["aspect_ratio"] == 2.0
    assert payload["palette"] == ["#112233"]


def test_build_creates_editable_pptx(tmp_path: Path) -> None:
    spec_path = tmp_path / "spec.json"
    fingerprint_path = tmp_path / "fingerprint.json"
    output_path = tmp_path / "output.pptx"
    spec_path.write_text(
        json.dumps(
            {
                "version": "1",
                "title": "Built",
                "slides": [{"layout": "title", "title": "Built", "subtitle": "CLI"}],
            }
        )
    )
    fingerprint_path.write_text(
        json.dumps(
            {
                "version": "1",
                "canvas": {"width": 10.0, "height": 5.0, "aspect_ratio": 2.0},
                "palette": ["#112233"],
                "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
                "geometry": {},
                "density": "balanced",
                "motif": "none",
            }
        )
    )

    result = run_cli(
        "build",
        str(spec_path),
        "--fingerprint",
        str(fingerprint_path),
        "--output",
        str(output_path),
    )

    assert result.returncode == 0, result.stderr
    assert output_path.read_bytes().startswith(b"PK")

    verification = run_cli("verify", str(output_path))

    assert verification.returncode == 0, verification.stderr
    assert json.loads(verification.stdout) == {
        "editable_text_shapes": 2,
        "slide_count": 1,
    }


def test_build_resolves_image_assets_relative_to_spec(tmp_path: Path) -> None:
    spec_dir = tmp_path / "project"
    assets = spec_dir / "assets"
    assets.mkdir(parents=True)
    Image.new("RGB", (4, 2), color=(17, 34, 51)).save(assets / "product.png")
    spec = spec_dir / "presentation.json"
    spec.write_text(
        json.dumps(
            {
                "version": "1",
                "title": "Product",
                "slides": [
                    {
                        "layout": "image_text",
                        "title": "Product",
                        "image": "assets/product.png",
                        "image_alt": "Product",
                        "body": ["Editable"],
                    }
                ],
            }
        )
    )
    output = tmp_path / "presentation.pptx"

    result = run_cli(
        "build",
        str(spec),
        "--fingerprint",
        "examples/style-fingerprint.json",
        "--output",
        str(output),
    )

    assert result.returncode == 0, result.stderr
    presentation = Presentation(str(output))
    assert any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for shape in presentation.slides[0].shapes
    )


def test_build_reports_stable_error_for_missing_image_asset(tmp_path: Path) -> None:
    spec = tmp_path / "presentation.json"
    spec.write_text(
        json.dumps(
            {
                "version": "1",
                "title": "Product",
                "slides": [
                    {
                        "layout": "image_text",
                        "title": "Product",
                        "image": "assets/missing.png",
                        "image_alt": "Missing",
                        "body": ["Editable"],
                    }
                ],
            }
        )
    )
    output = tmp_path / "presentation.pptx"

    result = run_cli(
        "build",
        str(spec),
        "--fingerprint",
        "examples/style-fingerprint.json",
        "--output",
        str(output),
    )

    assert result.returncode == 2
    assert result.stderr.startswith("INVALID_SPEC:")
    assert "assets/missing.png" in result.stderr
    assert "Traceback" not in result.stderr
    assert not output.exists()


def test_build_reports_stable_error_for_invalid_fingerprint(tmp_path: Path) -> None:
    spec_path = tmp_path / "spec.json"
    fingerprint_path = tmp_path / "fingerprint.json"
    output_path = tmp_path / "output.pptx"
    spec_path.write_text(
        json.dumps(
            {
                "version": "1",
                "title": "Built",
                "slides": [{"layout": "title", "title": "Built"}],
            }
        )
    )
    fingerprint_path.write_text(
        json.dumps(
            {
                "version": "1",
                "canvas": {"width": 10.0, "height": 5.0, "aspect_ratio": 2.0},
                "palette": [],
                "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
                "geometry": {},
                "density": "balanced",
                "motif": "none",
            }
        )
    )

    result = run_cli(
        "build",
        str(spec_path),
        "--fingerprint",
        str(fingerprint_path),
        "--output",
        str(output_path),
    )

    assert result.returncode == 2
    assert result.stderr.startswith("INVALID_SPEC:")
    assert "palette" in result.stderr
    assert not output_path.exists()


def test_build_reports_stable_error_when_output_parent_is_missing(tmp_path: Path) -> None:
    output = tmp_path / "missing" / "presentation.pptx"

    result = run_cli(
        "build",
        "examples/presentation.json",
        "--fingerprint",
        "examples/style-fingerprint.json",
        "--output",
        str(output),
    )

    assert result.returncode == 2
    assert result.stderr.startswith("OUTPUT_WRITE_FAILED:")
    assert str(output) in result.stderr
    assert "Traceback" not in result.stderr
    assert not output.exists()


def test_verify_reports_stable_error_for_invalid_package(tmp_path: Path) -> None:
    artifact = tmp_path / "broken.pptx"
    artifact.write_text("not a package")

    result = run_cli("verify", str(artifact))

    assert result.returncode != 0
    assert result.stderr.startswith("VERIFY_FAILED:")
    assert str(artifact) in result.stderr
