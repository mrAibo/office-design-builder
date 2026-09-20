from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from office_design_builder.inspectors.pptx import inspect_pptx


def test_inspect_pptx_extracts_dimensions_layouts_fonts_and_colors(
    tmp_path: Path,
) -> None:
    reference_path = tmp_path / "reference.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = text_box.text_frame.paragraphs[0].add_run()
    run.text = "Evidence"
    run.font.name = "Aptos"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    presentation.save(str(reference_path))

    fingerprint = inspect_pptx(reference_path)
    repeated = inspect_pptx(reference_path)

    assert fingerprint == repeated
    assert fingerprint.canvas == {"width": 10.0, "height": 5.0, "aspect_ratio": 2.0}
    assert fingerprint.typography == {
        "heading_font": "Aptos",
        "body_font": "Aptos",
    }
    assert fingerprint.palette == ["#112233"]
    assert fingerprint.geometry["slide_count"] == 1
    assert "Blank" in fingerprint.geometry["layout_names"]
