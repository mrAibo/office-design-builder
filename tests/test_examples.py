import json
from pathlib import Path
from time import sleep

from pptx import Presentation

from office_design_builder.models import PresentationSpecV1, StyleFingerprintV1
from office_design_builder.renderers.pptx import build_presentation


EXAMPLES = Path(__file__).parents[1] / "examples"


def test_nine_slide_example_builds_deterministically_with_native_objects(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        json.loads((EXAMPLES / "presentation.json").read_text())
    )
    fingerprint = StyleFingerprintV1.from_dict(
        json.loads((EXAMPLES / "style-fingerprint.json").read_text())
    )
    assert [slide["layout"] for slide in spec.slides] == [
        "title",
        "two_column",
        "section",
        "title_bullets",
        "image_text",
        "comparison",
        "timeline",
        "table",
        "chart",
    ]

    first = tmp_path / "first.pptx"
    second = tmp_path / "second.pptx"
    build_presentation(spec, fingerprint, first, asset_root=EXAMPLES)
    sleep(1.1)
    build_presentation(spec, fingerprint, second, asset_root=EXAMPLES)

    assert first.read_bytes() == second.read_bytes()
    presentation = Presentation(str(first))
    assert len(presentation.slides) == 9
    assert sum(shape.has_table for slide in presentation.slides for shape in slide.shapes) == 1
    assert sum(shape.has_chart for slide in presentation.slides for shape in slide.shapes) == 1
    assert sum(shape.shape_type == 13 for slide in presentation.slides for shape in slide.shapes) == 1
