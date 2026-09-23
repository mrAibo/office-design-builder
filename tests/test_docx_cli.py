"""End-to-end CLI contract for DOCX without external services."""
import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]


def run_cli(*args: object) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-m", "office_design_builder.cli", *(str(arg) for arg in args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
    )


def test_docx_example_build_verify_and_determinism(tmp_path: Path) -> None:
    spec = ROOT / "examples/document.json"
    fingerprint = ROOT / "examples/style-fingerprint.json"
    assert run_cli("validate", spec).returncode == 0
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    for output in (a, b):
        result = run_cli("build", spec, "--fingerprint", fingerprint, "--output", output)
        assert result.returncode == 0, result.stderr
    assert a.read_bytes() == b.read_bytes()
    checked = run_cli("verify", a)
    assert checked.returncode == 0, checked.stderr
    result = json.loads(checked.stdout)
    assert result["editable_paragraphs"] >= 4
    assert result["table_count"] == 1
    assert result["image_count"] == 1
    document = Document(str(a))
    assert [(p.style.name if p.style else "", p.text) for p in document.paragraphs if p.text] == [
        ("Heading 1", "Quarterly review"),
        ("Normal", "This document demonstrates editable native Word content."),
        ("Heading 2", "Highlights"),
        ("List Bullet", "Progress is visible"),
        ("List Bullet", "Next steps are actionable"),
    ]
    assert document.inline_shapes[0]._inline.docPr.get("descr") == "Example dashboard"
    assert document.tables[0].cell(1, 0).text == "Coverage"


@pytest.mark.parametrize("payload", [
    {"version": "1", "title": "x", "slides": [], "blocks": []},
    {"version": "1", "title": "x"},
])
def test_ambiguous_or_missing_spec_kind_fails(tmp_path: Path, payload: dict) -> None:
    spec = tmp_path / "invalid.json"
    spec.write_text(json.dumps(payload))
    result = run_cli("validate", spec)
    assert result.returncode == 2
    assert "INVALID_SPEC" in result.stderr


def test_build_rejects_wrong_output_extension(tmp_path: Path) -> None:
    result = run_cli(
        "build", ROOT / "examples/document.json", "--fingerprint",
        ROOT / "examples/style-fingerprint.json", "--output", tmp_path / "bad.pptx",
    )
    assert result.returncode == 2
    assert "INVALID_SPEC" in result.stderr
    assert not (tmp_path / "bad.pptx").exists()


def test_verify_corrupt_docx_is_handled(tmp_path: Path) -> None:
    artifact = tmp_path / "bad.docx"
    artifact.write_bytes(b"not a zip")
    result = run_cli("verify", artifact)
    assert result.returncode == 2
    assert "VERIFY_FAILED" in result.stderr


def test_symlink_image_escape_is_rejected_before_output(tmp_path: Path) -> None:
    outside = tmp_path / "outside.png"
    outside.write_bytes((ROOT / "examples/assets/dashboard.png").read_bytes())
    spec_dir = tmp_path / "spec"
    spec_dir.mkdir()
    try:
        (spec_dir / "linked.png").symlink_to(outside)
    except OSError as exc:
        pytest.skip(f"symlink creation unavailable on this runner: {exc}")
    spec = json.loads((ROOT / "examples/document.json").read_text())
    spec["blocks"][-1]["image"]["path"] = "linked.png"
    source = spec_dir / "document.json"
    source.write_text(json.dumps(spec))
    output = tmp_path / "escape.docx"
    result = run_cli("build", source, "--fingerprint", ROOT / "examples/style-fingerprint.json", "--output", output)
    assert result.returncode == 2, result.stderr
    assert "INVALID_SPEC" in result.stderr
    assert not output.exists()


def test_document_block_unknown_fields_rejected(tmp_path: Path) -> None:
    spec = {"version": "1", "title": "A", "blocks": [{"heading": {"text": "A", "level": 1, "typo": "x"}}]}
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(spec))
    result = run_cli("validate", path)
    assert result.returncode == 2
    assert "INVALID_SPEC" in result.stderr


def test_bool_heading_level_rejected(tmp_path: Path) -> None:
    spec = {"version": "1", "title": "A", "blocks": [{"heading": {"text": "A", "level": True}}]}
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(spec))
    result = run_cli("validate", path)
    assert result.returncode == 2
    assert "INVALID_SPEC" in result.stderr
