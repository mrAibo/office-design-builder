from pathlib import Path
from time import sleep
from typing import Any, cast

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from office_design_builder.models import PresentationSpecV1, StyleFingerprintV1
from office_design_builder.renderers.pptx import build_presentation


def test_build_presentation_creates_editable_title_and_two_column_slides(
    tmp_path: Path,
) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Quarterly Review",
            "slides": [
                {"layout": "title", "title": "Quarterly Review", "subtitle": "Q3"},
                {
                    "layout": "two_column",
                    "title": "Results",
                    "left": ["Revenue increased"],
                    "right": ["Costs decreased"],
                },
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 10.0, "height": 5.0, "aspect_ratio": 2.0},
            "palette": ["#112233", "#DDEEFF"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "none",
        }
    )
    output_path = tmp_path / "output.pptx"

    build_presentation(spec, fingerprint, output_path)

    result = Presentation(str(output_path))
    assert len(result.slides) == 2
    assert result.slide_width == 9_144_000
    assert result.slide_height == 4_572_000
    text_shapes = [
        cast(Any, shape)
        for slide in result.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    texts = [shape.text for shape in text_shapes if shape.text]
    assert texts == [
        "Quarterly Review",
        "Q3",
        "Results",
        "Revenue increased",
        "Costs decreased",
    ]
    assert all(
        shape.shape_type != MSO_SHAPE_TYPE.PICTURE
        for slide in result.slides
        for shape in slide.shapes
    )
    title_shape = next(shape for shape in text_shapes if shape.text == "Quarterly Review")
    title_run = title_shape.text_frame.paragraphs[0].runs[0]
    assert str(title_run.font.color.rgb) == "112233"
    assert title_run.font.name == "Aptos"


def test_build_presentation_uses_palette_accents_and_balanced_panels(
    tmp_path: Path,
) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Visual system",
            "slides": [
                {
                    "layout": "two_column",
                    "title": "Balanced content",
                    "left": ["Left"],
                    "right": ["Right"],
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "offset_blocks",
        }
    )
    output_path = tmp_path / "polished.pptx"

    build_presentation(spec, fingerprint, output_path)

    result = Presentation(str(output_path))
    slide = result.slides[0]
    shapes: dict[str, Any] = {shape.name: shape for shape in slide.shapes}
    slide_width = result.slide_width
    slide_height = result.slide_height
    assert slide_width is not None
    assert slide_height is not None
    assert {"Title accent", "Left panel", "Right panel"} <= shapes.keys()
    assert str(shapes["Title accent"].fill.fore_color.rgb) == "17324D"
    assert str(shapes["Left panel"].fill.fore_color.rgb) == "F9FBFC"
    assert shapes["Left panel"].width == shapes["Right panel"].width
    assert shapes["Left panel"].height <= slide_height * 0.25
    assert shapes["Left panel"].left < slide_width / 2
    assert shapes["Right panel"].left > slide_width / 2
    assert all(
        shape.left >= 0
        and shape.top >= 0
        and shape.left + shape.width <= slide_width
        and shape.top + shape.height <= slide_height
        for shape in slide.shapes
    )


def test_title_slide_uses_compact_decorative_motif(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Title",
            "slides": [{"layout": "title", "title": "Title", "subtitle": "Subtitle"}],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "offset_blocks",
        }
    )
    output_path = tmp_path / "title.pptx"

    build_presentation(spec, fingerprint, output_path)

    result = Presentation(str(output_path))
    shapes: dict[str, Any] = {shape.name: shape for shape in result.slides[0].shapes}
    slide_width = result.slide_width
    slide_height = result.slide_height
    assert slide_width is not None
    assert slide_height is not None
    assert "Hero block" not in shapes
    assert {"Motif back", "Motif front"} <= shapes.keys()
    assert shapes["Motif back"].width < slide_width * 0.12
    assert shapes["Motif front"].height < slide_height * 0.22


def test_build_presentation_renders_editable_section_slide(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Architecture",
            "slides": [
                {
                    "layout": "section",
                    "title": "Architecture",
                    "subtitle": "How the pieces fit together",
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "offset_blocks",
        }
    )
    output = tmp_path / "section.pptx"

    build_presentation(spec, fingerprint, output)

    result = Presentation(str(output))
    slide = result.slides[0]
    section_shapes = [cast(Any, shape) for shape in slide.shapes]
    shapes: dict[str, Any] = {shape.name: shape for shape in section_shapes}
    texts = [shape.text for shape in section_shapes if shape.has_text_frame and shape.text]
    slide_width = result.slide_width
    slide_height = result.slide_height
    assert slide_width is not None
    assert slide_height is not None
    assert texts == ["Architecture", "How the pieces fit together"]
    assert {"Section field", "Section accent"} <= shapes.keys()
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in section_shapes)
    assert all(
        shape.left >= 0
        and shape.top >= 0
        and shape.left + shape.width <= slide_width
        and shape.top + shape.height <= slide_height
        for shape in section_shapes
    )


def test_build_presentation_renders_native_title_bullets(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Priorities",
            "slides": [
                {
                    "layout": "title_bullets",
                    "title": "Priorities",
                    "bullets": ["Reliability", "Editability", "Determinism"],
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "offset_blocks",
        }
    )
    output = tmp_path / "title-bullets.pptx"

    build_presentation(spec, fingerprint, output)

    result = Presentation(str(output))
    slide = result.slides[0]
    text_shapes = [cast(Any, shape) for shape in slide.shapes if shape.has_text_frame]
    assert [shape.text for shape in text_shapes if shape.text] == [
        "Priorities",
        "Reliability\nEditability\nDeterminism",
    ]
    bullet_shape = next(shape for shape in text_shapes if shape.text.startswith("Reliability"))
    assert [paragraph.text for paragraph in bullet_shape.text_frame.paragraphs] == [
        "Reliability",
        "Editability",
        "Determinism",
    ]
    assert all("<a:buChar" in paragraph._p.xml for paragraph in bullet_shape.text_frame.paragraphs)


def test_build_presentation_embeds_image_text_asset(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (400, 200), color=(23, 50, 77)).save(assets / "product.png")
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Product",
            "slides": [
                {
                    "layout": "image_text",
                    "title": "Product",
                    "image": "assets/product.png",
                    "image_alt": "Product dashboard",
                    "body": ["Local-first workflow", "Editable output"],
                    "image_position": "left",
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "offset_blocks",
        }
    )
    output = tmp_path / "image-text.pptx"
    repeated = tmp_path / "image-text-repeated.pptx"

    build_presentation(spec, fingerprint, output, asset_root=tmp_path)
    build_presentation(spec, fingerprint, repeated, asset_root=tmp_path)

    assert output.read_bytes() == repeated.read_bytes()
    result = Presentation(str(output))
    slide = result.slides[0]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert pictures[0].name == "Image: Product dashboard"
    assert pictures[0].width / pictures[0].height == pytest.approx(2.0, rel=0.01)
    texts = [cast(Any, shape).text for shape in slide.shapes if shape.has_text_frame and cast(Any, shape).text]
    assert texts == ["Product", "Local-first workflow\nEditable output"]


def test_build_presentation_rejects_missing_image_without_output(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Product",
            "slides": [
                {
                    "layout": "image_text",
                    "title": "Product",
                    "image": "assets/missing.png",
                    "image_alt": "Missing",
                    "body": ["Body"],
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "none",
        }
    )
    output = tmp_path / "missing.pptx"

    with pytest.raises(ValueError, match="assets/missing.png"):
        build_presentation(spec, fingerprint, output, asset_root=tmp_path)

    assert not output.exists()


def test_build_presentation_renders_editable_comparison(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Build or buy",
            "slides": [
                {
                    "layout": "comparison",
                    "title": "Build or buy",
                    "left_title": "Build",
                    "left": ["Full control", "Higher effort"],
                    "right_title": "Buy",
                    "right": ["Fast adoption", "Vendor dependency"],
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "offset_blocks",
        }
    )
    output = tmp_path / "comparison.pptx"

    build_presentation(spec, fingerprint, output)

    result = Presentation(str(output))
    slide = result.slides[0]
    shapes = [cast(Any, shape) for shape in slide.shapes]
    names = {shape.name for shape in shapes}
    texts = [shape.text for shape in shapes if shape.has_text_frame and shape.text]
    assert {"Comparison left panel", "Comparison right panel"} <= names
    assert texts == [
        "Build or buy",
        "Build",
        "Full control\nHigher effort",
        "Buy",
        "Fast adoption\nVendor dependency",
    ]


def test_build_presentation_renders_editable_timeline(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Delivery plan",
            "slides": [
                {
                    "layout": "timeline",
                    "title": "Delivery plan",
                    "events": [
                        {"label": "Q1", "description": "Prototype"},
                        {"label": "Q2", "description": "Pilot"},
                        {"label": "Q3", "description": "Launch"},
                    ],
                }
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 13.333, "height": 7.5, "aspect_ratio": 1.7777},
            "palette": ["#17324D", "#E7EEF5"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "none",
        }
    )
    output_path = tmp_path / "timeline.pptx"

    build_presentation(spec, fingerprint, output_path)

    result = Presentation(str(output_path))
    slide = result.slides[0]
    text_shapes = [cast(Any, shape) for shape in slide.shapes if shape.has_text_frame]
    texts = [shape.text for shape in text_shapes if shape.text]
    assert texts == ["Delivery plan", "Q1", "Prototype", "Q2", "Pilot", "Q3", "Launch"]
    assert {"Timeline rail", "Timeline marker 1", "Timeline marker 2", "Timeline marker 3"} <= {
        shape.name for shape in slide.shapes
    }
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)


def test_build_presentation_is_byte_deterministic(tmp_path: Path) -> None:
    spec = PresentationSpecV1.from_dict(
        {
            "version": "1",
            "title": "Determinism",
            "slides": [
                {"layout": "title", "title": "Same", "subtitle": "Input"},
                {
                    "layout": "section",
                    "title": "Architecture",
                    "subtitle": "Deterministic section",
                },
                {
                    "layout": "title_bullets",
                    "title": "Priorities",
                    "bullets": ["Reliable", "Editable"],
                },
            ],
        }
    )
    fingerprint = StyleFingerprintV1.from_dict(
        {
            "version": "1",
            "canvas": {"width": 10.0, "height": 5.0, "aspect_ratio": 2.0},
            "palette": ["#112233"],
            "typography": {"heading_font": "Aptos", "body_font": "Aptos"},
            "geometry": {},
            "density": "balanced",
            "motif": "none",
        }
    )
    first = tmp_path / "first.pptx"
    second = tmp_path / "second.pptx"

    build_presentation(spec, fingerprint, first)
    sleep(2.1)
    build_presentation(spec, fingerprint, second)

    assert first.read_bytes() == second.read_bytes()
