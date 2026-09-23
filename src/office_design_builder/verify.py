"""Structural verification for generated Office artifacts."""

from __future__ import annotations

from pathlib import Path
from docx import Document

from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError

from office_design_builder.errors import VerifyError


def verify_pptx(path: Path) -> dict[str, int]:
    try:
        presentation = Presentation(str(path))
    except (OSError, ValueError, KeyError, PptxPackageNotFoundError) as exc:
        raise VerifyError(f"{path}: not a readable PPTX package") from exc

    editable_text_shapes = sum(
        1
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame and getattr(shape, "text", "").strip()
    )
    if not presentation.slides:
        raise VerifyError(f"{path}: presentation contains no slides")
    if editable_text_shapes == 0:
        raise VerifyError(f"{path}: presentation contains no editable text shapes")
    return {
        "slide_count": len(presentation.slides),
        "editable_text_shapes": editable_text_shapes,
    }


def verify_docx(path: Path) -> dict[str, int]:
    try:
        document = Document(str(path))
    except (OSError, ValueError, KeyError, TypeError, RuntimeError, DocxPackageNotFoundError) as exc:
        raise VerifyError(f"{path}: not a readable DOCX package") from exc
    paragraphs = sum(1 for paragraph in document.paragraphs if paragraph.text.strip())
    tables = len(document.tables)
    images = len(document.inline_shapes)
    if not paragraphs and not tables:
        raise VerifyError(f"{path}: document contains no editable content")
    return {"editable_paragraphs": paragraphs, "table_count": tables, "image_count": images}