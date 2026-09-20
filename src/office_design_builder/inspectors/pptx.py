"""Deterministic style evidence extraction from PowerPoint packages."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE

from office_design_builder.models import StyleFingerprintV1

_EMU_PER_INCH = 914_400


def inspect_pptx(path: Path) -> StyleFingerprintV1:
    presentation = Presentation(str(path))
    if presentation.slide_width is None or presentation.slide_height is None:
        raise ValueError("PPTX slide dimensions are unavailable")
    width = presentation.slide_width / _EMU_PER_INCH
    height = presentation.slide_height / _EMU_PER_INCH

    fonts: set[str] = set()
    colors: set[str] = set()
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text_shape = cast(Any, shape)
            for paragraph in text_shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
                    if run.font.color.type == MSO_COLOR_TYPE.RGB:
                        colors.add(f"#{run.font.color.rgb}")

    font = sorted(fonts)[0] if fonts else "sans-serif"
    layout_names = sorted({layout.name for layout in presentation.slide_layouts})
    return StyleFingerprintV1(
        version="1",
        canvas={"width": width, "height": height, "aspect_ratio": width / height},
        palette=sorted(colors),
        typography={"heading_font": font, "body_font": font},
        geometry={"slide_count": len(presentation.slides), "layout_names": layout_names},
        density="balanced",
        motif="pptx",
    )
