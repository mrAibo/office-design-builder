"""Deterministic editable DOCX rendering."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile, ZipInfo

from PIL import Image
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import RGBColor

from office_design_builder.errors import InvalidSpecError, InputNotFoundError
from office_design_builder.models import DocumentSpecV1, StyleFingerprintV1

_FIXED_TIMESTAMP = datetime(2000, 1, 1)
_FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)


def _canonicalize_zip_bytes(source: bytes) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(source), "r") as package, ZipFile(output, "w") as canonical_package:
        for original in package.infolist():
            data = package.read(original.filename)
            if original.filename.endswith(".xlsx"):
                data = _canonicalize_zip_bytes(data)
            if original.filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)>).*?(</dcterms:(?:created|modified)>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                    data,
                )
            canonical = ZipInfo(original.filename, date_time=_FIXED_ZIP_DATE)
            canonical.compress_type = original.compress_type
            canonical.external_attr = original.external_attr
            canonical.internal_attr = original.internal_attr
            canonical.create_system = original.create_system
            canonical_package.writestr(canonical, data)
    return output.getvalue()


def _canonicalize_package(path: Path) -> None:
    canonical = _canonicalize_zip_bytes(path.read_bytes())
    temporary = path.with_suffix(".tmp")
    temporary.write_bytes(canonical)
    temporary.replace(path)


def _tint(color: RGBColor, ratio: float = 0.25) -> RGBColor:
    return RGBColor(
        *(round(255 + (channel - 255) * ratio) for channel in color)
    )


def _add_heading(document: Document, text: str, level: int, font_name: str, font_size: int, color: RGBColor, bold: bool = True) -> None:
    style_name = f"Heading {level}"
    style = document.styles[style_name]
    style.font.name = font_name
    style.font.size = Pt(font_size)
    style.font.bold = bold
    style.font.color.rgb = color
    document.add_paragraph(text, style=style_name)


def _add_paragraph_text(document: Document, text: str, font_name: str, font_size: int, color: RGBColor, align: WD_PARAGRAPH_ALIGNMENT | None = None) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = color

    if align is not None:
        paragraph.alignment = align


def _add_bullets(document: Document, items: list[str], font_name: str, font_size: int, color: RGBColor) -> None:
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.style = document.styles["List Bullet"]
        run = paragraph.add_run(item)
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        paragraph.paragraph_format.left_indent = Inches(0.5)


def _add_table(document: Document, columns: list[str], rows: list[list[str]], font_name: str, font_size: int, color: RGBColor) -> None:
    # Add one row for header, then the data rows
    table = document.add_table(rows=len(rows) + 1, cols=len(columns))
    table.style = document.styles["Table Grid"]

    # Add header row
    for col_index, column_title in enumerate(columns):
        cell = table.cell(0, col_index)
        cell.text = column_title
        cell.paragraphs[0].runs[0].font.name = font_name
        cell.paragraphs[0].runs[0].font.size = Pt(font_size)
        cell.paragraphs[0].runs[0].font.color.rgb = color
        cell.paragraphs[0].runs[0].font.bold = True

    # Add data rows
    for row_index, row_cells in enumerate(rows, start=1):
        for col_index, cell_text in enumerate(row_cells):
            cell = table.cell(row_index, col_index)
            cell.text = cell_text
            cell.paragraphs[0].runs[0].font.name = font_name
            cell.paragraphs[0].runs[0].font.size = Pt(font_size)
            cell.paragraphs[0].runs[0].font.color.rgb = color


def _add_image(document: Document, image_path: Path, alt_text: str, max_width: float) -> None:
    if not image_path.exists():
        raise InputNotFoundError(f"Image not found: {image_path}")

    try:
        with Image.open(image_path) as img:
            img.verify()
    except (OSError, ValueError) as exc:
        raise InvalidSpecError(f"Invalid image asset: {image_path}") from exc

    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(max_width))
    inline_shape._inline.docPr.set("descr", alt_text)


def _preflight_images(spec: DocumentSpecV1, asset_root: Path | None) -> dict[str, tuple[Path, int, int]]:
    image_blocks = [block for block in spec.blocks if list(block.keys())[0] == "image"]
    if not image_blocks:
        return {}

    if asset_root is None:
        raise InvalidSpecError("asset_root is required for image blocks")

    root = asset_root.resolve()
    assets: dict[str, tuple[Path, int, int]] = {}

    for block in image_blocks:
        block_type = list(block.keys())[0]
        if block_type == "image":
            image_info = block["image"]
            relative = image_info["path"].replace("\\", "/")
            path = (root / relative).resolve()

            if path != root and root not in path.parents:
                raise InvalidSpecError(f"image asset escapes asset_root: {image_info['path']}")

            if not path.exists():
                raise InputNotFoundError(f"Image not found: {path}")

            try:
                with Image.open(path) as image:
                    width, height = image.size
                    image.verify()
            except (OSError, ValueError) as exc:
                raise InvalidSpecError(f"Invalid image asset: {image_info['path']}") from exc

            assets[image_info["path"]] = (path, width, height)

    return assets


def build_document(
    spec: DocumentSpecV1,
    fingerprint: StyleFingerprintV1,
    output_path: Path,
    *,
    asset_root: Path | None = None,
) -> Path:
    assets = _preflight_images(spec, asset_root)
    document = Document()
    for section in document.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    for block in spec.blocks:
        block_type = list(block.keys())[0]

        if block_type == "heading":
            heading_info = block["heading"]
            text = heading_info["text"]
            level = heading_info.get("level", 1)
            color = RGBColor(
                int(fingerprint.palette[0][1:3], 16),
                int(fingerprint.palette[0][3:5], 16),
                int(fingerprint.palette[0][5:7], 16)
            )
            font_name = fingerprint.typography["heading_font"]
            font_size = 24 if level == 1 else 18 if level == 2 else 14

            _add_heading(document, text, level, font_name, font_size, color)

        elif block_type == "paragraph":
            paragraph_info = block["paragraph"]
            text = paragraph_info["text"]
            color = RGBColor(
                int(fingerprint.palette[0][1:3], 16),
                int(fingerprint.palette[0][3:5], 16),
                int(fingerprint.palette[0][5:7], 16)
            )
            font_name = fingerprint.typography["body_font"]
            font_size = 11

            _add_paragraph_text(document, text, font_name, font_size, color)

        elif block_type == "bullets":
            bullets_info = block["bullets"]
            items = bullets_info["items"]
            color = RGBColor(
                int(fingerprint.palette[0][1:3], 16),
                int(fingerprint.palette[0][3:5], 16),
                int(fingerprint.palette[0][5:7], 16)
            )
            font_name = fingerprint.typography["body_font"]
            font_size = 11

            _add_bullets(document, items, font_name, font_size, color)

        elif block_type == "table":
            table_info = block["table"]
            columns = table_info["columns"]
            rows = table_info["rows"]
            color = RGBColor(
                int(fingerprint.palette[0][1:3], 16),
                int(fingerprint.palette[0][3:5], 16),
                int(fingerprint.palette[0][5:7], 16)
            )
            font_name = fingerprint.typography["body_font"]
            font_size = 10

            _add_table(document, columns, rows, font_name, font_size, color)

        elif block_type == "image":
            image_info = block["image"]
            image_path = assets[image_info["path"]][0]
            alt_text = image_info["alt"]
            _add_image(document, image_path, alt_text, max_width=6.0)

    document.core_properties.created = _FIXED_TIMESTAMP
    document.core_properties.modified = _FIXED_TIMESTAMP
    document.core_properties.title = spec.title

    document.save(str(output_path))

    _canonicalize_package(output_path)

    return output_path