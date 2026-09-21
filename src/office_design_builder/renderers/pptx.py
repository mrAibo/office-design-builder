"""Deterministic editable PPTX rendering."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import ZipFile, ZipInfo

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from office_design_builder.errors import InvalidSpecError
from office_design_builder.models import PresentationSpecV1, StyleFingerprintV1

_FIXED_TIMESTAMP = datetime(2000, 1, 1)
_FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)


def _canonicalize_package(path: Path) -> None:
    with ZipFile(path, "r") as package:
        entries = [(info, package.read(info.filename)) for info in package.infolist()]

    temporary = path.with_suffix(".tmp")
    with ZipFile(temporary, "w") as package:
        for original, data in entries:
            canonical = ZipInfo(original.filename, date_time=_FIXED_ZIP_DATE)
            canonical.compress_type = original.compress_type
            canonical.external_attr = original.external_attr
            canonical.internal_attr = original.internal_attr
            canonical.create_system = original.create_system
            package.writestr(canonical, data)
    temporary.replace(path)


def _tint(color: RGBColor, ratio: float = 0.25) -> RGBColor:
    return RGBColor(
        *(round(255 + (channel - 255) * ratio) for channel in color)
    )


def _preflight_images(
    spec: PresentationSpecV1, asset_root: Path | None
) -> dict[str, tuple[Path, int, int]]:
    image_slides = [slide for slide in spec.slides if slide["layout"] == "image_text"]
    if not image_slides:
        return {}
    if asset_root is None:
        raise InvalidSpecError("asset_root is required for image_text slides")
    root = asset_root.resolve()
    assets: dict[str, tuple[Path, int, int]] = {}
    for slide in image_slides:
        relative = slide["image"].replace("\\", "/")
        path = (root / relative).resolve()
        if path != root and root not in path.parents:
            raise InvalidSpecError(f"image asset escapes asset_root: {slide['image']}")
        try:
            with Image.open(path) as image:
                width, height = image.size
                image.verify()
        except (OSError, ValueError) as exc:
            raise InvalidSpecError(f"invalid image asset: {slide['image']}") from exc
        assets[slide["image"]] = (path, width, height)
    return assets


def _add_panel(
    slide: Any,
    name: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    color: RGBColor,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_text(
    slide: Any,
    text: str,
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    font_name: str,
    font_size: int,
    color: RGBColor,
    bold: bool = False,
) -> None:
    shape = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.text_frame.margin_left = 0
    shape.text_frame.margin_right = 0
    shape.text_frame.margin_top = 0
    shape.text_frame.margin_bottom = 0
    paragraph = shape.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def _add_bullets(
    slide: Any,
    items: list[str],
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    font_name: str,
    font_size: int,
    color: RGBColor,
) -> None:
    shape = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.name = "Bullet list"
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        properties = paragraph._p.get_or_add_pPr()
        properties.set("marL", str(Inches(0.3)))
        properties.set("indent", str(-Inches(0.18)))
        properties.append(bullet)
        paragraph.space_after = Pt(10)
        run = paragraph.add_run()
        run.text = item
        run.font.name = font_name
        run.font.size = Pt(font_size)
        run.font.color.rgb = color


def build_presentation(
    spec: PresentationSpecV1,
    fingerprint: StyleFingerprintV1,
    output_path: Path,
    *,
    asset_root: Path | None = None,
) -> Path:
    image_assets = _preflight_images(spec, asset_root)
    presentation = Presentation()
    presentation.core_properties.created = _FIXED_TIMESTAMP
    presentation.core_properties.modified = _FIXED_TIMESTAMP
    presentation.core_properties.last_modified_by = "office-design-builder"
    presentation.core_properties.revision = 1
    presentation.slide_width = Inches(fingerprint.canvas["width"])
    presentation.slide_height = Inches(fingerprint.canvas["height"])
    blank_layout = presentation.slide_layouts[6]
    primary = RGBColor.from_string(fingerprint.palette[0].removeprefix("#"))
    secondary_source = fingerprint.palette[1] if len(fingerprint.palette) > 1 else fingerprint.palette[0]
    secondary = RGBColor.from_string(secondary_source.removeprefix("#"))
    panel_color = _tint(secondary)
    heading_font = fingerprint.typography["heading_font"]
    body_font = fingerprint.typography["body_font"]
    canvas_width = fingerprint.canvas["width"]
    canvas_height = fingerprint.canvas["height"]
    margin = canvas_width * 0.06
    heading_size = 36 if canvas_width >= 12 else 30
    body_size = 18 if canvas_width >= 12 else 16

    for slide_spec in spec.slides:
        slide = presentation.slides.add_slide(blank_layout)
        if slide_spec["layout"] == "title":
            _add_panel(
                slide,
                "Title accent",
                left=margin,
                top=canvas_height * 0.23,
                width=0.12,
                height=canvas_height * 0.32,
                color=primary,
            )
            _add_panel(
                slide,
                "Motif back",
                left=canvas_width * 0.82,
                top=canvas_height * 0.34,
                width=canvas_width * 0.09,
                height=canvas_height * 0.18,
                color=panel_color,
            )
            _add_panel(
                slide,
                "Motif front",
                left=canvas_width * 0.85,
                top=canvas_height * 0.41,
                width=canvas_width * 0.09,
                height=canvas_height * 0.18,
                color=primary,
            )
            _add_text(
                slide,
                slide_spec["title"],
                left=margin + 0.32,
                top=canvas_height * 0.29,
                width=canvas_width * 0.62,
                height=canvas_height * 0.14,
                font_name=heading_font,
                font_size=heading_size,
                color=primary,
                bold=True,
            )
            _add_text(
                slide,
                slide_spec.get("subtitle", ""),
                left=margin + 0.32,
                top=canvas_height * 0.48,
                width=canvas_width * 0.62,
                height=canvas_height * 0.1,
                font_name=body_font,
                font_size=body_size,
                color=primary,
            )
            continue

        if slide_spec["layout"] == "section":
            field_left = margin
            field_top = canvas_height * 0.18
            field_width = canvas_width - 2 * margin
            field_height = canvas_height * 0.64
            _add_panel(
                slide,
                "Section field",
                left=field_left,
                top=field_top,
                width=field_width,
                height=field_height,
                color=panel_color,
            )
            _add_panel(
                slide,
                "Section accent",
                left=field_left,
                top=field_top,
                width=0.16,
                height=field_height,
                color=primary,
            )
            _add_text(
                slide,
                slide_spec["title"],
                left=field_left + 0.55,
                top=canvas_height * 0.36,
                width=field_width * 0.72,
                height=canvas_height * 0.16,
                font_name=heading_font,
                font_size=heading_size + 4,
                color=primary,
                bold=True,
            )
            if "subtitle" in slide_spec:
                _add_text(
                    slide,
                    slide_spec["subtitle"],
                    left=field_left + 0.55,
                    top=canvas_height * 0.56,
                    width=field_width * 0.72,
                    height=canvas_height * 0.1,
                    font_name=body_font,
                    font_size=body_size,
                    color=primary,
                )
            continue

        if slide_spec["layout"] == "title_bullets":
            _add_panel(
                slide,
                "Title accent",
                left=margin,
                top=canvas_height * 0.085,
                width=0.12,
                height=canvas_height * 0.095,
                color=primary,
            )
            _add_text(
                slide,
                slide_spec["title"],
                left=margin + 0.32,
                top=canvas_height * 0.08,
                width=canvas_width - 2 * margin - 0.32,
                height=canvas_height * 0.12,
                font_name=heading_font,
                font_size=heading_size - 6,
                color=primary,
                bold=True,
            )
            _add_bullets(
                slide,
                slide_spec["bullets"],
                left=margin + 0.25,
                top=canvas_height * 0.29,
                width=canvas_width - 2 * margin - 0.5,
                height=canvas_height * 0.55,
                font_name=body_font,
                font_size=body_size + 2,
                color=primary,
            )
            continue

        if slide_spec["layout"] == "image_text":
            _add_panel(
                slide,
                "Title accent",
                left=margin,
                top=canvas_height * 0.085,
                width=0.12,
                height=canvas_height * 0.095,
                color=primary,
            )
            _add_text(
                slide,
                slide_spec["title"],
                left=margin + 0.32,
                top=canvas_height * 0.08,
                width=canvas_width - 2 * margin - 0.32,
                height=canvas_height * 0.12,
                font_name=heading_font,
                font_size=heading_size - 6,
                color=primary,
                bold=True,
            )
            gap = canvas_width * 0.045
            region_width = (canvas_width - 2 * margin - gap) / 2
            region_top = canvas_height * 0.27
            region_height = canvas_height * 0.58
            image_left = margin
            text_left = margin + region_width + gap
            if slide_spec.get("image_position", "left") == "right":
                image_left, text_left = text_left, image_left
            path, pixel_width, pixel_height = image_assets[slide_spec["image"]]
            scale = min(region_width / pixel_width, region_height / pixel_height)
            picture_width = pixel_width * scale
            picture_height = pixel_height * scale
            picture = slide.shapes.add_picture(
                str(path),
                Inches(image_left + (region_width - picture_width) / 2),
                Inches(region_top + (region_height - picture_height) / 2),
                width=Inches(picture_width),
                height=Inches(picture_height),
            )
            picture.name = f"Image: {slide_spec['image_alt']}"
            body_height = 0.55 * len(slide_spec["body"])
            body_top = region_top + (region_height - body_height) / 2
            _add_bullets(
                slide,
                slide_spec["body"],
                left=text_left + 0.15,
                top=body_top,
                width=region_width - 0.3,
                height=body_height,
                font_name=body_font,
                font_size=body_size,
                color=primary,
            )
            continue

        if slide_spec["layout"] == "comparison":
            gap = canvas_width * 0.035
            panel_width = (canvas_width - 2 * margin - gap) / 2
            panel_top = canvas_height * 0.27
            panel_height = canvas_height * 0.58
            right_left = margin + panel_width + gap
            _add_panel(
                slide,
                "Title accent",
                left=margin,
                top=canvas_height * 0.085,
                width=0.12,
                height=canvas_height * 0.095,
                color=primary,
            )
            _add_text(
                slide,
                slide_spec["title"],
                left=margin + 0.32,
                top=canvas_height * 0.08,
                width=canvas_width - 2 * margin - 0.32,
                height=canvas_height * 0.12,
                font_name=heading_font,
                font_size=heading_size - 6,
                color=primary,
                bold=True,
            )
            for side, panel_left in (("left", margin), ("right", right_left)):
                _add_panel(
                    slide,
                    f"Comparison {side} panel",
                    left=panel_left,
                    top=panel_top,
                    width=panel_width,
                    height=panel_height,
                    color=panel_color,
                )
                _add_text(
                    slide,
                    slide_spec[f"{side}_title"],
                    left=panel_left + 0.35,
                    top=panel_top + 0.3,
                    width=panel_width - 0.7,
                    height=0.45,
                    font_name=heading_font,
                    font_size=body_size + 4,
                    color=primary,
                    bold=True,
                )
                _add_bullets(
                    slide,
                    slide_spec[side],
                    left=panel_left + 0.25,
                    top=panel_top + 1.0,
                    width=panel_width - 0.5,
                    height=panel_height - 1.3,
                    font_name=body_font,
                    font_size=body_size,
                    color=primary,
                )
            continue

        gap = canvas_width * 0.035
        item_count = max(len(slide_spec["left"]), len(slide_spec["right"]))
        panel_height = min(canvas_height * 0.42, 0.7 + 0.5 * item_count)
        panel_top = min(canvas_height * 0.32, canvas_height - panel_height - 0.55)
        panel_width = (canvas_width - 2 * margin - gap) / 2
        right_left = margin + panel_width + gap
        _add_panel(
            slide,
            "Title accent",
            left=margin,
            top=canvas_height * 0.085,
            width=0.12,
            height=canvas_height * 0.095,
            color=primary,
        )
        _add_panel(
            slide,
            "Left panel",
            left=margin,
            top=panel_top,
            width=panel_width,
            height=panel_height,
            color=panel_color,
        )
        _add_panel(
            slide,
            "Right panel",
            left=right_left,
            top=panel_top,
            width=panel_width,
            height=panel_height,
            color=panel_color,
        )
        _add_text(
            slide,
            slide_spec["title"],
            left=margin + 0.32,
            top=canvas_height * 0.08,
            width=canvas_width - 2 * margin - 0.32,
            height=canvas_height * 0.12,
            font_name=heading_font,
            font_size=heading_size - 6,
            color=primary,
            bold=True,
        )
        _add_text(
            slide,
            "\n".join(slide_spec["left"]),
            left=margin + 0.32,
            top=panel_top + 0.35,
            width=panel_width - 0.64,
            height=panel_height - 0.7,
            font_name=body_font,
            font_size=body_size,
            color=primary,
        )
        _add_text(
            slide,
            "\n".join(slide_spec["right"]),
            left=right_left + 0.32,
            top=panel_top + 0.35,
            width=panel_width - 0.64,
            height=panel_height - 0.7,
            font_name=body_font,
            font_size=body_size,
            color=primary,
        )

    presentation.save(str(output_path))
    _canonicalize_package(output_path)
    return output_path
